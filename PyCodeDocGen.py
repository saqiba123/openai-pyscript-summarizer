import ast
import os
import logging
import openai
from docx import Document
from fpdf import FPDF
from dotenv import load_dotenv

# Configure logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

# Load environment variables
load_dotenv()

# OpenAI API setup
openai_key = os.getenv("OPENAI_API_KEY")
if not openai_key:
    logger.error("OpenAI API key not found. Please set OPENAI_API_KEY in .env file.")
    raise ValueError("Missing OpenAI API key")

client = openai.OpenAI(api_key=openai_key)

def generate_detailed_explanation(code_snippet: str) -> str:
    """
    Generate a detailed, line-by-line explanation of a code snippet using OpenAI API.
    
    Args:
        code_snippet (str): The code snippet to explain
    
    Returns:
        str: Detailed explanation of the code
    """
    try:
        response = client.chat.completions.create(
            model="gpt-3.5-turbo",
            messages=[
                {
                    "role": "system", 
                    "content": "You are an expert Python code explainer. Provide a detailed, line-by-line explanation of the code snippet, breaking down its purpose, functionality, and any important nuances."
                },
                {
                    "role": "user", 
                    "content": f"Please provide a comprehensive, line-by-line explanation of this code:\n\n{code_snippet}"
                }
            ],
            max_tokens=500
        )
        return response.choices[0].message.content.strip()
    except Exception as e:
        logger.error(f"Error generating explanation: {e}")
        return f"Unable to generate explanation. Error: {str(e)}"

def extract_script_info(file_path: str):
    """
    Extract detailed information about the Python script.
    
    Args:
        file_path (str): Path to the Python script
    
    Returns:
        tuple: Extracted script information
    """
    try:
        with open(file_path, "r", encoding="utf-8") as file:
            script_content = file.read()
            tree = ast.parse(script_content, filename=file_path)
    except Exception as e:
        logger.error(f"Error parsing script: {e}")
        return [], [], [], []

    imports = []
    functions = []
    classes = []
    other_code = []

    # Tracking lines
    lines = script_content.split('\n')
    code_lines = set(range(len(lines)))

    # Detailed AST traversal
    for node in ast.walk(tree):
        if isinstance(node, ast.Import):
            imports.extend([alias.name for alias in node.names])
            code_lines.discard(node.lineno - 1)
        elif isinstance(node, ast.ImportFrom):
            imports.append(node.module)
            code_lines.discard(node.lineno - 1)
        elif isinstance(node, ast.FunctionDef):
            # Extract function details with explanation
            func_code = ''.join(lines[node.lineno-1:node.end_lineno])
            func_info = {
                "name": node.name,
                "params": [arg.arg for arg in node.args.args],
                "code": func_code,
                "explanation": generate_detailed_explanation(func_code)
            }
            functions.append(func_info)
            code_lines.discard(node.lineno - 1)
        elif isinstance(node, ast.ClassDef):
            # Extract class details with explanation
            class_code = ''.join(lines[node.lineno-1:node.end_lineno])
            methods = [n.name for n in node.body if isinstance(n, ast.FunctionDef)]
            class_info = {
                "name": node.name,
                "methods": methods,
                "code": class_code,
                "explanation": generate_detailed_explanation(class_code)
            }
            classes.append(class_info)
            code_lines.discard(node.lineno - 1)

    # Capture remaining code with explanation
    other_code_lines = [lines[i].strip() for i in code_lines if lines[i].strip()]
    for code_line in other_code_lines:
        other_code.append({
            "code": code_line,
            "explanation": generate_detailed_explanation(code_line)
        })

    return imports, functions, classes, other_code

def generate_pdf_document(file_path: str):
    """
    Generate a detailed PDF document for the Python script.
    
    Args:
        file_path (str): Path to the Python script
    """
    # Extract script information
    imports, functions, classes, other_code = extract_script_info(file_path)
    script_name = os.path.basename(file_path)
    output_file = f"{os.path.splitext(script_name)[0]}_detailed_summary.pdf"

    # Create PDF
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)

    # Title
    pdf.set_font("Arial", "B", 16)
    pdf.cell(0, 10, f"Detailed Script Analysis: {script_name}", ln=True, align="C")
    pdf.ln(10)

    # Imports
    pdf.set_font("Arial", "B", 12)
    pdf.cell(0, 10, "Imported Libraries:", ln=True)
    pdf.set_font("Arial", size=10)
    for imp in imports:
        pdf.cell(0, 10, f"- {imp}", ln=True)
    pdf.ln(5)

    # Classes
    pdf.set_font("Arial", "B", 12)
    pdf.cell(0, 10, "Classes:", ln=True)
    for cls in classes:
        pdf.set_font("Arial", "B", 11)
        pdf.cell(0, 10, f"Class: {cls['name']}", ln=True)
        pdf.set_font("Arial", size=10)
        pdf.cell(0, 10, f"Methods: {', '.join(cls['methods'])}", ln=True)
        
        # Code and Explanation
        pdf.set_font("Courier", size=9)
        pdf.multi_cell(0, 5, f"Code:\n{cls['code']}")
        pdf.set_font("Arial", size=10)
        pdf.multi_cell(0, 5, f"Explanation:\n{cls['explanation']}")
        pdf.ln(5)

    # Functions
    pdf.set_font("Arial", "B", 12)
    pdf.cell(0, 10, "Functions:", ln=True)
    for func in functions:
        pdf.set_font("Arial", "B", 11)
        pdf.cell(0, 10, f"Function: {func['name']}", ln=True)
        pdf.set_font("Arial", size=10)
        pdf.cell(0, 10, f"Parameters: {', '.join(func['params'])}", ln=True)
        
        # Code and Explanation
        pdf.set_font("Courier", size=9)
        pdf.multi_cell(0, 5, f"Code:\n{func['code']}")
        pdf.set_font("Arial", size=10)
        pdf.multi_cell(0, 5, f"Explanation:\n{func['explanation']}")
        pdf.ln(5)

    # Other Code
    pdf.set_font("Arial", "B", 12)
    pdf.cell(0, 10, "Other Code:", ln=True)
    for code_block in other_code:
        pdf.set_font("Courier", size=9)
        pdf.multi_cell(0, 5, f"Code:\n{code_block['code']}")
        pdf.set_font("Arial", size=10)
        pdf.multi_cell(0, 5, f"Explanation:\n{code_block['explanation']}")
        pdf.ln(5)

    # Save PDF
    pdf.output(output_file)
    logger.info(f"Detailed documentation saved as {output_file}")

# Example usage
generate_pdf_document("bank_test.py")